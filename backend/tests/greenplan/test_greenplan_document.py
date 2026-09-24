"""greenplan/document.py -- пояснительная записка GreenPlan в DOCX."""

import io

import pytest
from docx import Document

from greenplan.document import MAX_VIOLATION_ROWS, build_document, project_title, provenance
from greenplan.pattern_assignment import ZoneAssignment
from helpers import make_object, make_scene, make_zone


def _assignment(**overrides) -> ZoneAssignment:
    fields = dict(
        zone_id="path_corridor_001", zone_kind="path_corridor", pattern_id="linear_hedge_row",
        source_project="02_peschany_pereulok", source_quote="кусты с шагом 1.4 м -- вдоль дороги",
        confidence=0.67, zone_area_sqm=120.0,
        tree_species=["Клен остролистный"], bush_species=["Спирея японская"],
        species_basis="ассортимент для озеленения Москвы, категория «магистрали и улицы»",
    )
    fields.update(overrides)
    return ZoneAssignment(**fields)


def _generated(obj_id, obj_type, x, z, species, catalog_id):
    return make_object(obj_id, obj_type, x, z, metadata={
        "generated": True, "species": species, "catalogId": catalog_id, "zone_id": "path_corridor_001",
    })


def _scene():
    building = make_zone(type="building")  # [-5,5]x[-5,5]
    objects = [
        _generated("tree_new_1", "tree", 0, 30, "Клен остролистный", "species_klen_ostrolistnyy"),
        _generated("bush_new_1", "bush", 20, 0, "Спирея японская", "species_spireya_yaponskaya"),
        _generated("bush_new_2", "bush", 22, 0, "Спирея японская", "species_spireya_yaponskaya"),
        # Существующее дерево в 1 м от стены -- нарушение исходного состояния.
        make_object("tree_old_1", "tree", 6, 0),
    ]
    return make_scene(restrictions=[building], objects=objects)


def _text(docx_bytes: bytes) -> str:
    doc = Document(io.BytesIO(docx_bytes))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def test_document_has_all_sections():
    text = _text(build_document(_scene(), [_assignment()], title="Тестовый участок"))
    for heading in (
        "Пояснительная записка к проекту озеленения",
        "1. Общие сведения об участке",
        "2. Нормативная база",
        "3. Принятые решения",
        "4. Ведомость элементов озеленения",
        "5. Проверка соответствия нормам",
        "6. Технические требования к посадке",
        "7. Ограничения и допущения",
    ):
        assert heading in text
    assert "Тестовый участок" in text


def test_decision_carries_its_provenance_and_species():
    text = _text(build_document(_scene(), [_assignment()]))
    assert "по аналогии с проектом «Песчаный переулок»" in text
    assert "Спирея японская" in text
    assert "--" not in text  # ASCII-тире из подписей в коде заменены на «—»


def test_fallback_decision_is_marked_for_review():
    fallback = _assignment(source_project=None, source_quote=None, confidence=0.0)
    assert "проверить" in provenance(fallback)
    assert "проверить" in _text(build_document(_scene(), [fallback]))


def test_schedule_follows_gost_form_9_and_counts_only_new_plants():
    doc = Document(io.BytesIO(build_document(_scene(), [_assignment()])))
    schedule = next(t for t in doc.tables if t.rows[0].cells[1].text == "Наименование породы или вида насаждения")
    assert [c.text for c in schedule.rows[0].cells] == [
        "Поз.", "Наименование породы или вида насаждения", "Возраст, лет", "Кол.", "Примечание",
    ]
    rows = {row.cells[1].text: row.cells[3].text for row in schedule.rows[1:]}
    lawn = rows.pop("Газон обыкновенный")
    assert rows == {"Клен остролистный": "1", "Спирея японская": "2"}
    # Деревья -- первыми, как в примере формы, газон -- последним и в м².
    assert schedule.rows[1].cells[1].text == "Клен остролистный"
    assert schedule.rows[-1].cells[1].text == "Газон обыкновенный"
    assert lawn.endswith(" м²") and int(lawn.removesuffix(" м²").replace(" ", "")) > 9000
    assert "травосмеси" in schedule.rows[-1].cells[4].text


def test_new_lawn_brings_seeding_and_watering_requirements():
    text = _text(build_document(_scene(), [_assignment()]))
    assert "устройство нового" in text
    assert "п. 9.24: норма высева газонных трав" in text
    assert "полив газона — 10 л/м²" in text
    assert "Газон допускается над инженерными сетями" in text


def test_lawn_is_recomputed_not_taken_from_request():
    # Сцена пришла без газона (старый клиент или ручные правки после
    # GreenPlan) -- записка всё равно считает газон по текущим посадкам.
    scene = _scene()
    assert scene.lawns == []
    assert "Газон обыкновенный" in _text(build_document(scene, [_assignment()]))


def test_violations_split_new_and_existing():
    text = _text(build_document(_scene(), [_assignment()]))
    assert "Нарушений у новых посадок: 0" in text
    assert "Нарушений у существующих объектов: 1" in text
    assert "Существующие объекты с нарушениями" in text


def test_new_plant_violations_are_listed_individually():
    scene = _scene()
    scene.objects.append(_generated("tree_new_bad", "tree", 7, 0, "Клен остролистный", "species_klen_ostrolistnyy"))
    text = _text(build_document(scene, [_assignment()]))
    assert "Нарушения у новых посадок" in text
    assert "tree_new_bad" in text
    assert MAX_VIOLATION_ROWS >= 1


def test_numbers_use_decimal_comma():
    text = _text(build_document(_scene(), [_assignment()]))
    assert "1,5" in text  # кустарник от здания
    assert "| 1.5 |" not in text


def test_ai_text_only_as_marked_appendix():
    without = _text(build_document(_scene(), [_assignment()]))
    assert "Приложение А" not in without
    with_report = _text(build_document(_scene(), [_assignment()], report="Вдоль дорожек -- живая изгородь."))
    assert "Приложение А. Обоснование решений (текст ИИ — проверить)" in with_report
    assert "Вдоль дорожек" in with_report


def test_empty_assignments_still_produce_a_document():
    text = _text(build_document(make_scene(), []))
    assert "Зон, пригодных для посадки, на участке нет." in text


@pytest.mark.parametrize(
    ("slug", "title"),
    [
        ("10_stary_gay", "улица Старый Гай"),
        ("28_formal_bosque_primer", "синтетический эталон «формальный боскет»"),
        ("no_such_project", "no_such_project"),
    ],
)
def test_project_title_from_rationale(slug, title):
    assert project_title(slug) == title
