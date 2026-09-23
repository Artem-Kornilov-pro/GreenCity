"""Общие фикстуры сквозных тестов: настоящее приложение (main.app) через HTTP
(TestClient), без подмен внутри бэкенда. Подменены только внешние границы:
MongoDB/Redis (in-memory, фикстуры fake_mongo/fake_redis из tests/conftest.py)
и клиенты LLM -- облачный для правки текстом и локальный Ollama для отчёта.
Всё остальное -- парсер, GreenPlan, нормы, экспорт, документ -- работает
по-настоящему, на реальных DXF из locations/."""

import io
import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from fastapi.testclient import TestClient

from core.paths import LOCATIONS_DIR
from greenplan import decision_report
from main import app
from text_editor import llm_client


def pytest_collection_modifyitems(items):
    """Все тесты этой папки -- с маркером e2e (make test-e2e / pytest -m e2e)."""
    here = Path(__file__).parent
    for item in items:
        if here in Path(item.fspath).parents:
            item.add_marker(pytest.mark.e2e)


@pytest.fixture
def client(fake_mongo, fake_redis):
    with TestClient(app) as c:
        yield c


def location_dxf(slug: str) -> Path:
    for path in (LOCATIONS_DIR / slug / f"{slug}.dxf", *LOCATIONS_DIR.glob(f"location_old/{slug}/*.dxf")):
        if path.exists():
            return path
    raise FileNotFoundError(slug)


def upload_dxf(client: TestClient, path: Path, name: str | None = None) -> dict:
    r = client.post("/api/parse", files={"file": (name or path.name, path.read_bytes(), "application/dxf")})
    assert r.status_code == 200, r.text
    return r.json()


def docx_text(content: bytes) -> str:
    doc = Document(io.BytesIO(content))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


class FakeChatLLM:
    """OpenAI-совместимый клиент Chat Completions: отдаёт заранее заданный
    ответ и запоминает, что ему прислали (проверяем, что в промпт ушли
    настоящие факты сцены)."""

    def __init__(self, content: str):
        self.content = content
        self.requests: list[dict] = []
        self.chat = SimpleNamespace(completions=self)

    def create(self, **kwargs):
        self.requests.append(kwargs)
        return SimpleNamespace(
            choices=[SimpleNamespace(message=SimpleNamespace(content=self.content), finish_reason="stop")],
            usage=SimpleNamespace(prompt_tokens=100, completion_tokens=50),
        )

    def prompt_text(self) -> str:
        return "\n".join(m["content"] for req in self.requests for m in req["messages"])


@pytest.fixture
def fake_text_editor_llm(monkeypatch):
    """Подменить облачную LLM правки текстом; план задаёт тест через
    fake.content = json.dumps({...})."""
    monkeypatch.setenv("LLM_PROVIDER", "gemini")
    fake = FakeChatLLM(json.dumps({"operations": [], "explanation": ""}))
    monkeypatch.setattr(llm_client, "_client_and_model", lambda: (fake, "fake-model"))
    return fake


@pytest.fixture
def fake_report_llm(monkeypatch):
    """Подменить локальную LLM (Ollama) текста-отчёта GreenPlan."""
    fake = FakeChatLLM("Вдоль дорожек участка предложена живая изгородь, открытая площадь занята группами деревьев.")
    monkeypatch.setattr(decision_report, "_client", lambda: fake)
    return fake


@pytest.fixture
def e2e():
    """Помощники сквозных тестов (conftest нельзя импортировать из теста)."""
    return SimpleNamespace(location_dxf=location_dxf, upload_dxf=upload_dxf, docx_text=docx_text)
