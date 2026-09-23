"""
Пояснительная записка GreenPlan в DOCX -- "сценарий выгрузки/генерации
документации для внедрения предложенного решения" из ТЗ кейса и Этап 6
issue #23 ("отчёт о решениях ... с пометкой происхождения каждого выбора").

Документ собирается детерминированно из того, что уже посчитано:
расстановка и решения по зонам (deterministic_placement.generate_for_scene),
проверка норм (violation_report), ведомость (assortment_report), нормы и
выписки из актов (setback_norms, data/norms/). Единственная часть от LLM --
текст-обоснование из /api/greenplan/report, если фронтенд его передал; он
вынесен в приложение и помечен как требующий проверки.

У каждого решения -- происхождение, как требует issue #23:
    "по нормативу"                 -- отступы и допустимость видов: посчитано
                                      по актам, проверять не нужно;
    "по аналогии с проектом X"     -- паттерн взят из похожего прошлого проекта;
    "типовое решение -- проверить" -- ни один похожий проект такой зоны не
                                      размечал, паттерн по умолчанию;
    "текст ИИ -- проверить"        -- приложение с пересказом от LLM.

Ведомость -- по форме 9 ГОСТ 21.508-2020 ("Ведомость элементов озеленения").
DOCX, а не PDF: проектировщик правит записку перед выпуском, а PDF из Word
печатается в один шаг.
"""

from __future__ import annotations

import io
from collections import Counter
from datetime import date
from functools import lru_cache
from typing import Optional

import yaml
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from core.paths import LOCATIONS_DIR, NORMS_DIR
from core.plant_catalog import catalog_by_id
from core.schemas import Scene
from core.setback_norms import SETBACK_NORMS, species_rules
from greenplan.assortment_report import AssortmentRow, summarize_assortment
from greenplan.decision_report import ZONE_KIND_LABELS
from greenplan.pattern_assignment import ZoneAssignment
from greenplan.pattern_library import PATTERN_LIBRARY
from greenplan.site_characterization import SiteCharacteristics, characterize_site
from greenplan.species_selection import in_base_515
from greenplan.violation_report import Violation, find_violations

_SP82_YAML = NORMS_DIR / "sp-82-13330-2016" / "planting.yaml"
_LOCATIONS = LOCATIONS_DIR

# Сколько строк нарушений выводить таблицей -- на реальных участках у
# СУЩЕСТВУЮЩИХ объектов их бывают сотни; остаток -- одной строкой-итогом.
MAX_VIOLATION_ROWS = 100

ZONE_TYPE_LABELS: dict[str, str] = {
    "building": "здания",
    "road": "проезжая часть",
    "pedestrian_path": "пешеходные дорожки и тротуары",
    "gas_pipeline": "газопровод",
    "sewer": "канализация, водосток, дренаж",
    "water_pipeline": "водопровод",
    "electrical": "силовой кабель",
    "signal_cable": "кабель связи",
    "heat_network": "теплосеть",
    "playground_zone": "детские площадки",
    "overhead_power_line": "воздушная ЛЭП",
    "transformer": "трансформаторная подстанция",
    "protected_zone": "охраняемая зона",
    "custom": "прочие зоны (парковки, неуточнённые сети)",
}

TERRITORY_LABELS: dict[str, str] = {
    "двор": "дворовая территория",
    "улица": "улица, проезд",
    "площадь": "площадь, общественное пространство",
    "парк_сквер": "парк, сквер, бульвар",
    "промышленная_охранная": "производственная или охранная зона",
    "неопределено": "не определён",
}

SETBACK_SOURCE = "СП 42.13330.2016, табл. 9.1; ППМ 743-ПП, табл. 3.6.1"

NORMATIVE_BASIS: list[tuple[str, str]] = [
    ("СП 42.13330.2016, п. 9.6, табл. 9.1", "Расстояния от зданий, сооружений и сетей до деревьев и кустарников"),
    ("ППМ 743-ПП, табл. 3.6.1", "Те же расстояния, в т.ч. «теплопровод, трубопровод, теплосеть» — 2 м; "
     "деревья с широкой кроной — не ближе 10 м от здания (прим. 3)"),
    ("МГСН 1.02-02 (ППМ 623-ПП), п. 4.2.8", "Породы, чувствительные к прогреву почвы, у теплотрасс"),
    ("СП 82.13330.2016, п. 9.22", "Колючие растения — не ближе 2 м от площадок и пешеходных коммуникаций"),
    ("ППМ 369-ПП, приложение 1", "Перечень инвазивных видов: ни один не включён в посадку"),
    ("Ассортимент для озеленения Москвы (основной, дополнительный, перспективный)",
     "Подбор видов по типу территории и кодам ограничений"),
    ("ППМ 515-ПП, табл. 4", "Базовый ассортимент — приоритет при подборе видов для дворовых территорий"),
    ("ГОСТ 21.508-2020, форма 9", "Форма ведомости элементов озеленения"),
]

LIMITATIONS: list[str] = [
    "Этап решений по существующим деревьям (сохранить, пересадить, вырубить) и расчёт компенсационной "
    "стоимости не выполнялись: в исходных данных нет перечётной ведомости.",
    "Назначение зданий в исходном чертеже не размечено, поэтому отступ 10 м для деревьев с широкой кроной "
    "применён ко всем зданиям, а отступ 10 м от школ и детских садов (743-ПП) не применялся.",
    "Не проверялись: инсоляция помещений (СП 42, прим. 3 к табл. 9.1), увеличение отступов для деревьев "
    "с кроной более 5 м (прим. 1), свободная высота 2,1 м над пешеходными путями (СП 59.13330.2020, п. 5.1.7).",
    "Отступы от инженерных сетей отсчитаны от края их охранного коридора, а не от оси сети, — то есть с "
    "запасом относительно нормы.",
    "Возраст и параметры посадочного материала в ведомости указаны только для видов базового ассортимента "
    "ППМ 515-ПП; для остальных уточняются на стадии рабочей документации.",
]


# --- Данные ---------------------------------------------------------------


@lru_cache(maxsize=64)
def project_title(slug: str) -> str:
    """Человекочитаемое название проекта-аналога из заголовка его
    design_rationale.md ("# Пояснительная записка ... — Песчаный переулок")."""
    path = _LOCATIONS / slug / "design_rationale.md"
    try:
        first = path.read_text(encoding="utf-8").splitlines()[0].lstrip("# ").strip()
    except (OSError, IndexError):
        return slug
    name = first.split("—", 1)[-1].strip() if "—" in first else first
    if first.startswith("Синтетический эталон"):
        return f"синтетический эталон «{name}»"
    return name


@lru_cache(maxsize=1)
def _sp82() -> dict:
    return yaml.safe_load(_SP82_YAML.read_text(encoding="utf-8"))


def provenance(assignment: ZoneAssignment) -> str:
    if assignment.source_project:
        return f"по аналогии с проектом «{project_title(assignment.source_project)}»"
    return "типовое решение (похожие проекты такую зону не размечали) — проверить"


class DecisionGroup:
    """Решения, сведённые по (вид зоны, паттерн, проект-источник) -- как и в
    decision_report._summarize: не "зона №142", а "вдоль дорожек -- изгородь"."""

    def __init__(self, members: list[ZoneAssignment]):
        self.members = members
        self.first = max(members, key=lambda a: a.confidence)

    @property
    def area(self) -> float:
        return sum(a.zone_area_sqm for a in self.members)


def _groups(assignments: list[ZoneAssignment]) -> list[DecisionGroup]:
    buckets: dict[tuple, list[ZoneAssignment]] = {}
    for a in assignments:
        buckets.setdefault((a.zone_kind, a.pattern_id, a.source_project), []).append(a)
    return sorted((DecisionGroup(m) for m in buckets.values()), key=lambda g: -g.area)


def _without_generated(scene: Scene) -> Scene:
    return scene.model_copy(update={"objects": [o for o in scene.objects if not o.metadata.get("generated")]})


# --- Оформление ------------------------------------------------------------


def _new_document():
    doc = Document()
    section = doc.sections[0]
    section.page_height, section.page_width = Mm(297), Mm(210)
    section.left_margin, section.right_margin = Mm(25), Mm(15)
    section.top_margin, section.bottom_margin = Mm(20), Mm(20)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    # Заголовки -- чёрным и тем же шрифтом: синий цвет темы Word по умолчанию
    # в пояснительной записке неуместен.
    for name, size in (("Title", 18), ("Heading 1", 14), ("Heading 2", 12)):
        font = doc.styles[name].font
        font.name = "Times New Roman"
        font.size = Pt(size)
        font.bold = True
        font.color.rgb = RGBColor(0, 0, 0)
    return doc


def _set_borders(table) -> None:
    """Явные границы всех ячеек -- стиль "Table Grid" рисуют не все
    просмотрщики (Quick Look, часть веб-просмотрщиков его игнорирует)."""
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), "000000")
        borders.append(element)
    table._tbl.tblPr.append(borders)


def _shade(cell, fill: str = "E7E6E6") -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def _repeat_header(row) -> None:
    flag = OxmlElement("w:tblHeader")
    flag.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(flag)


def _table(doc, header: list[str], rows: list[list[str]], widths_mm: Optional[list[float]] = None):
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    _set_borders(table)
    _repeat_header(table.rows[0])
    for cell, text in zip(table.rows[0].cells, header):
        cell.text = text
        _shade(cell)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = text
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(10)
    if widths_mm:
        table.autofit = False
        for column, width in zip(table.columns, widths_mm):
            column.width = Mm(width)
        for row in table.rows:
            for cell, width in zip(row.cells, widths_mm):
                cell.width = Mm(width)
    doc.add_paragraph()  # отбивка после таблицы
    return table


def _fmt_area(value: float) -> str:
    return f"{value:,.0f}".replace(",", " ")


def _text(value: str) -> str:
    """Подписи паттернов и цитаты в коде пишутся с "--" вместо тире."""
    return " ".join(value.split()).replace(" -- ", " — ").replace("--", "—")


def _num(value: float) -> str:
    """1.5 -> "1,5": десятичная запятая, как в документах на русском."""
    return f"{value:g}".replace(".", ",")


# --- Разделы ---------------------------------------------------------------


def _section_general(doc, title: str, c: Optional[SiteCharacteristics], scene: Scene, new_count: int) -> None:
    doc.add_heading("1. Общие сведения об участке", level=1)
    if c is None:
        doc.add_paragraph("Граница участка в исходных данных не найдена — характеристики не рассчитывались.")
        return
    territory = TERRITORY_LABELS.get(c.territory_type, c.territory_type)
    if c.territory_type_is_heuristic:
        territory += " (определён по геометрии участка; тип не подтверждён)"
    present = sorted({z.type for z in scene.restrictions if z.severity != "allowed"})
    networks = ", ".join(ZONE_TYPE_LABELS.get(t, t) for t in present) or "нет"
    existing = Counter(o.type for o in scene.objects if o.type in ("tree", "bush"))
    rows = [
        ["Объект", title],
        ["Площадь участка", f"{_fmt_area(c.total_area_sqm)} м²"],
        ["Площадь, пригодная для посадки деревьев", f"{_fmt_area(c.plantable_area_sqm)} м² ({c.plantable_ratio:.0%})"],
        ["Тип территории", territory],
        ["Зданий", str(c.building_count)],
        ["Ограничения на участке", networks],
        ["Существующих деревьев / кустарников", f"{existing.get('tree', 0)} / {existing.get('bush', 0)}"],
        ["Предложено новых посадок", str(new_count)],
    ]
    _table(doc, ["Показатель", "Значение"], rows, [70, 100])


def _section_norms(doc, scene: Scene, assignments: list[ZoneAssignment]) -> None:
    doc.add_heading("2. Нормативная база", level=1)
    _table(doc, ["Документ", "Что применено"], [list(r) for r in NORMATIVE_BASIS], [70, 100])

    doc.add_heading("2.1. Отступы посадок, применённые на участке", level=2)
    present = sorted({z.type for z in scene.restrictions if z.severity != "allowed"})
    rows = []
    for zone_type in present:
        norm = SETBACK_NORMS.get(zone_type)
        if norm:
            rows.append([ZONE_TYPE_LABELS.get(zone_type, zone_type), _num(norm["tree"]), _num(norm["bush"]), SETBACK_SOURCE])
        else:
            rows.append([ZONE_TYPE_LABELS.get(zone_type, zone_type), "—", "—", "табличной нормы нет — отступ по ширине самой зоны"])
    _table(doc, ["Ограничение", "Дерево, м", "Кустарник, м", "Источник"], rows, [55, 22, 25, 68])
    doc.add_paragraph(
        "Отступ отсчитывается от края зоны ограничения; для сетей зона уже включает охранный коридор, "
        "поэтому фактическое расстояние до оси сети больше нормы."
    )

    used = sorted({s for a in assignments for s in (*a.tree_species, *a.bush_species)})
    species_rows = []
    for name in used:
        for rule in sorted(species_rules(name), key=lambda r: r.zone_type):
            if rule.zone_type in present:
                species_rows.append([name, ZONE_TYPE_LABELS.get(rule.zone_type, rule.zone_type), _num(rule.distance_m), rule.source])
    if species_rows:
        doc.add_heading("2.2. Увеличенные отступы по породе", level=2)
        _table(doc, ["Вид", "От чего", "Не ближе, м", "Источник"], species_rows, [50, 45, 22, 53])


def _section_decisions(doc, assignments: list[ZoneAssignment]) -> None:
    doc.add_heading("3. Принятые решения", level=1)
    doc.add_paragraph(
        "Участок разбит на геометрические зоны; для каждой зоны приём озеленения выбран по аналогии с "
        "похожими реализованными проектами (поиск ближайших по признакам участка), виды растений — по "
        "ассортименту и нормам. Расстановка рассчитана детерминированно, с проверкой отступов в каждой точке."
    )
    if not assignments:
        doc.add_paragraph("Зон, пригодных для посадки, на участке нет.")
        return
    rows = []
    for g in _groups(assignments):
        a = g.first
        species = "; ".join(filter(None, [
            ("деревья: " + ", ".join(a.tree_species)) if a.tree_species else "",
            ("кустарники: " + ", ".join(a.bush_species)) if a.bush_species else "",
        ])) or "—"
        rows.append([
            ZONE_KIND_LABELS.get(a.zone_kind, a.zone_kind),
            _text(PATTERN_LIBRARY[a.pattern_id].label),
            f"{len(g.members)} / {_fmt_area(g.area)}",
            provenance(a),
            species,
        ])
    _table(doc, ["Место", "Приём", "Зон / м²", "Происхождение решения", "Виды (по нормативу)"], rows, [30, 38, 20, 40, 42])

    quotes = [(g.first.source_project, g.first.source_quote) for g in _groups(assignments) if g.first.source_quote]
    if quotes:
        doc.add_heading("3.1. Проекты-аналоги", level=2)
        for slug, quote in dict(quotes).items():
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(f"{project_title(slug)}: ").bold = True
            p.add_run(_text(quote).strip('"'))

    bases = sorted({a.species_basis for a in assignments if a.species_basis})
    if bases:
        doc.add_heading("3.2. Основание подбора видов", level=2)
        for basis in bases:
            doc.add_paragraph(_text(basis[0].upper() + basis[1:]) + ".", style="List Bullet")


def _section_schedule(doc, assortment: list[AssortmentRow]) -> None:
    doc.add_heading("4. Ведомость элементов озеленения", level=1)
    doc.add_paragraph("Форма 9 по ГОСТ 21.508-2020.")
    if not assortment:
        doc.add_paragraph("Новых посадок нет.")
        return
    ordered = sorted(assortment, key=lambda r: (r.category != "дерево", -r.count, r.species))
    # Примечание: параметры посадочного материала (высота, ком) есть только у
    # видов базового ассортимента 515-ПП (data/norms/515-pp/assortment_base.csv).
    base_515 = {item.label for item in catalog_by_id().values() if in_base_515(item)}
    rows = [
        [str(i), r.species, "—", str(r.count), "по ППМ 515-ПП, табл. 4" if r.species in base_515 else "—"]
        for i, r in enumerate(ordered, 1)
    ]
    _table(doc, ["Поз.", "Наименование породы или вида насаждения", "Возраст, лет", "Кол.", "Примечание"], rows, [12, 80, 20, 15, 43])
    trees = sum(r.count for r in assortment if r.category == "дерево")
    bushes = sum(r.count for r in assortment if r.category == "кустарник")
    doc.add_paragraph(f"Итого: деревьев — {trees}, кустарников — {bushes}.")


def _section_violations(doc, violations: list[Violation], generated_ids: set[str]) -> None:
    doc.add_heading("5. Проверка соответствия нормам", level=1)
    new = [v for v in violations if v.object_id in generated_ids]
    old = [v for v in violations if v.object_id not in generated_ids]
    doc.add_paragraph(
        f"Все объекты участка повторно проверены на отступы независимо от алгоритма расстановки. "
        f"Нарушений у новых посадок: {len({v.object_id for v in new})}. "
        f"Нарушений у существующих объектов: {len({v.object_id for v in old})}."
    )
    if new:
        doc.add_heading("5.1. Нарушения у новых посадок", level=2)
        rows = [
            [v.object_id, _text(v.message), _num(v.distance_m), _num(v.required_m), _severity(v.severity)]
            for v in new[:MAX_VIOLATION_ROWS]
        ]
        _table(doc, ["Объект", "Ограничение", "Расстояние, м", "Требуется, м", "Уровень"], rows, [26, 70, 22, 22, 30])
        if len(new) > MAX_VIOLATION_ROWS:
            doc.add_paragraph(f"…и ещё {len(new) - MAX_VIOLATION_ROWS} (полный список — в интерфейсе GreenPlan).")
    if old:
        # Существующие посадки -- не решение GreenPlan, а исходное состояние
        # участка; не скрываем, но сводим по видам ограничений, а не сотнями строк.
        doc.add_heading("5.2. Существующие объекты с нарушениями (исходное состояние участка)", level=2)
        by_zone: dict[tuple[str, str], list[Violation]] = {}
        for v in old:
            by_zone.setdefault((ZONE_TYPE_LABELS.get(v.zone_type, v.zone_type), v.severity), []).append(v)
        rows = [
            [label, str(len({v.object_id for v in group})), _num(min(v.distance_m for v in group)),
             _num(max(v.required_m for v in group)), _severity(severity)]
            for (label, severity), group in sorted(by_zone.items(), key=lambda item: -len(item[1]))
        ]
        _table(doc, ["Ограничение", "Объектов", "Мин. расстояние, м", "Требуется, м", "Уровень"], rows, [60, 22, 30, 28, 30])


def _severity(value: str) -> str:
    return "запрет" if value == "forbidden" else "предупреждение"


def _section_technical(doc, has_hedges: bool) -> None:
    doc.add_heading("6. Технические требования к посадке", level=1)
    data = _sp82()
    doc.add_paragraph("По СП 82.13330.2016 «Благоустройство территорий».")
    items = [f"п. {r['clause']}: {' '.join(r['text'].split())}" for r in data["planting_pits"] if r["clause"] in ("9.6", "9.7", "9.9")]
    if has_hedges:
        items += [f"п. {r['clause']}: {' '.join(r['text'].split())}" for r in data["planting_pits"] if r["clause"] == "7.3"]
    w = data["watering_litres"]
    items.append(
        f"п. {w['clause']}: полив при посадке — {w['standard_sapling']} л на стандартный саженец, "
        f"{w['tree_root_ball_up_to_1x1_m']} л на дерево с комом до 1×1 м, {w['tree_root_ball_1x1_m_and_more']} л — "
        f"с комом 1×1 м и более, {w['bush']} л на куст."
    )
    items += [f"п. {data['acceptance']['clause']}: {text}." for text in data["acceptance"]["items"]]
    for text in items:
        doc.add_paragraph(_text(text), style="List Bullet")


def _section_limitations(doc, assignments: list[ZoneAssignment], c: Optional[SiteCharacteristics]) -> None:
    doc.add_heading("7. Ограничения и допущения", level=1)
    items = list(LIMITATIONS)
    if c is not None and c.territory_type == "неопределено":
        items.insert(0, "Тип территории по геометрии участка не определён — ассортимент подобран как для дворовых территорий.")
    if any(a.species_basis and "вне ассортимента" in a.species_basis for a in assignments):
        items.insert(0, "Для части зон в каталоге не нашлось видов из ассортимента Москвы — использован весь каталог, проверить.")
    for text in items:
        doc.add_paragraph(text, style="List Bullet")


def _section_ai_text(doc, report: Optional[str]) -> None:
    if not report:
        return
    doc.add_heading("Приложение А. Обоснование решений (текст ИИ — проверить)", level=1)
    doc.add_paragraph(
        "Текст ниже сформирован языковой моделью по списку решений раздела 3. Модели запрещено добавлять "
        "факты, которых нет в этом списке, но перед выпуском текст нужно сверить с разделом 3."
    ).runs[0].italic = True
    for paragraph in report.split("\n"):
        if paragraph.strip():
            doc.add_paragraph(paragraph.strip())


# --- Сборка ----------------------------------------------------------------


def build_document(scene: Scene, assignments: list[ZoneAssignment], report: Optional[str] = None, title: Optional[str] = None) -> bytes:
    """DOCX пояснительной записки по уже посчитанной сцене GreenPlan (с
    новыми объектами, metadata.generated) и её решениям по зонам."""
    title = (title or "").strip() or "Участок озеленения"
    base = _without_generated(scene)
    characteristics = characterize_site(base)
    generated = [o for o in scene.objects if o.metadata.get("generated")]
    violations = find_violations(scene)
    assortment = summarize_assortment(generated, catalog_by_id())

    doc = _new_document()
    heading = doc.add_heading("Пояснительная записка к проекту озеленения", level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph(title)
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta = doc.add_paragraph(f"Сформировано сервисом GreenCity (GreenPlan) {date.today():%d.%m.%Y}")
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.runs[0].italic = True

    _section_general(doc, title, characteristics, base, len(generated))
    _section_norms(doc, base, assignments)
    _section_decisions(doc, assignments)
    _section_schedule(doc, assortment)
    _section_violations(doc, violations, {o.id for o in generated})
    _section_technical(doc, has_hedges=any(a.pattern_id in ("linear_hedge_row", "building_ring") for a in assignments))
    _section_limitations(doc, assignments, characteristics)
    _section_ai_text(doc, report)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
